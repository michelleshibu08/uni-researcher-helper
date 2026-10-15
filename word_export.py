"""Turn a reference list into a Word document, formatted the way essays expect."""

from io import BytesIO

from docx import Document
from docx.shared import Inches, Pt

# What each style calls its reference list.
HEADINGS = {
    "apa": "References",
    "harvard": "Reference list",
    "mla": "Works Cited",
    "chicago": "Bibliography",
    "ieee": "References",
    "oscola": "Bibliography",
}


def new_document():
    document = Document()
    font = document.styles["Normal"].font
    font.name = "Times New Roman"
    font.size = Pt(12)
    return document


def add_reference_list(document, style, references):
    document.add_heading(HEADINGS.get(style, "References"), level=1)
    for reference in references:
        paragraph = document.add_paragraph(reference)
        # A "hanging indent": the first line starts at the margin and the rest are
        # indented, which is how APA, Harvard, MLA and Chicago lay out reference lists.
        paragraph.paragraph_format.left_indent = Inches(0.5)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5)
        paragraph.paragraph_format.space_after = Pt(12)


def to_bytes(document):
    # Save into memory instead of a file on disk, so it can be sent straight to the browser.
    file = BytesIO()
    document.save(file)
    return file.getvalue()


def reference_list_docx(style, references):
    """Return the bytes of a .docx file with a heading and one paragraph per reference."""
    document = new_document()
    add_reference_list(document, style, references)
    return to_bytes(document)


def essay_docx(title, draft, style, references):
    """Return the bytes of a .docx file with the essay, then its reference list on a new page."""
    document = new_document()
    document.add_heading(title, level=0)
    # Each line the student wrote becomes a paragraph, with the 1.5 line spacing many unis ask for.
    for text in draft.splitlines():
        if text.strip():
            paragraph = document.add_paragraph(text.strip())
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(10)
    if references:
        document.add_page_break()
        add_reference_list(document, style, references)
    return to_bytes(document)
